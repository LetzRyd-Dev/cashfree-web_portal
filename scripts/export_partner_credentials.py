import os
import pandas as pd
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.database import engine
from sqlalchemy import text

def generate_credentials_excel():
    downloads_path = os.path.expanduser("~/Downloads")
    output_file = os.path.join(downloads_path, "LetzRyd_Partner_Login_Credentials.xlsx")
    
    print(f"Connecting to database and fetching clean credentials...")
    with engine.connect() as conn:
        # 1. Independent Drivers (Login Enabled)
        ind_drivers_query = text("""
            SELECT 
                d.app_driver_id AS "App Driver ID",
                COALESCE(
                    NULLIF(d.driver_code, ''), 
                    'LR-DRV-' || LPAD(d.app_driver_id::text, 4, '0')
                ) AS "Driver Code (LetzRyd)",
                d.full_name AS "Driver Name",
                d.phone AS "Registered Mobile",
                '1234 / SMS OTP' AS "Default OTP",
                COALESCE(d.vehicle_reg_number, 'Unassigned') AS "Assigned Vehicle",
                COALESCE(d.vehicle_make, '') AS "Vehicle Make",
                COALESCE(d.vehicle_model, '') AS "Vehicle Model",
                'Independent Driver (Direct)' AS "Account Type",
                'LetzRyd Direct Operations' AS "Managing Entity",
                COALESCE(d.cw_to_collect, 0) AS "Current Outstanding (₹)",
                CASE WHEN d.is_active THEN 'Active (Can Login)' ELSE 'Inactive' END AS "Login Status"
            FROM app_drivers d
            WHERE d.operator_id IS NULL OR d.operator_id = 0
            ORDER BY d.app_driver_id ASC;
        """)
        df_ind_drivers = pd.read_sql(ind_drivers_query, conn)

        # 2. Fleet Operators (Login Enabled)
        operators_query = text("""
            SELECT 
                o.app_operator_id AS "App Operator ID",
                COALESCE(
                    NULLIF(o.operator_code, ''), 
                    'LR-OPR-' || LPAD(o.app_operator_id::text, 4, '0')
                ) AS "Operator Code (LetzRyd)",
                o.company_name AS "Fleet / Company Name",
                COALESCE(o.contact_person_name, o.company_name) AS "Contact Person",
                o.phone AS "Registered Mobile",
                '1234 / SMS OTP' AS "Default OTP",
                'Fleet Operator' AS "Account Type",
                COALESCE(o.total_vehicles, 0) AS "Total Fleet Vehicles",
                COALESCE(o.active_vehicles, 0) AS "Active Vehicles",
                COALESCE(o.total_drivers, 0) AS "Total Fleet Drivers",
                COALESCE(o.cw_to_collect, 0) AS "Fleet Outstanding Due (₹)",
                CASE WHEN o.is_active THEN 'Active (Can Login)' ELSE 'Inactive' END AS "Login Status"
            FROM app_operators o
            ORDER BY o.app_operator_id ASC;
        """)
        df_operators = pd.read_sql(operators_query, conn)

        # 3. Fleet Managed Drivers (Login Disabled - Managed by Operator)
        fleet_drivers_query = text("""
            SELECT DISTINCT ON (d.app_driver_id)
                d.app_driver_id AS "App Driver ID",
                COALESCE(
                    NULLIF(d.driver_code, ''), 
                    'LR-DRV-' || LPAD(d.app_driver_id::text, 4, '0')
                ) AS "Driver Code (LetzRyd)",
                d.full_name AS "Driver Name",
                d.phone AS "Registered Mobile",
                'LOGIN DISABLED' AS "Default OTP",
                COALESCE(o.company_name, 'Fleet Operator #' || d.operator_id) AS "Fleet Owner Name",
                COALESCE(d.vehicle_reg_number, 'Fleet Assigned') AS "Assigned Vehicle",
                COALESCE(d.vehicle_make, '') AS "Vehicle Make",
                COALESCE(d.vehicle_model, '') AS "Vehicle Model",
                'Fleet Driver (Managed)' AS "Account Type",
                'Managed by Fleet Operator' AS "Access Note",
                COALESCE(d.cw_to_collect, 0) AS "Driver Period Balance (₹)",
                'Disabled (Operator Portal Only)' AS "Login Status"
            FROM app_drivers d
            LEFT JOIN app_operators o ON d.operator_id = o.app_operator_id
            WHERE d.operator_id IS NOT NULL AND d.operator_id > 0
            ORDER BY d.app_driver_id ASC;
        """)
        df_fleet_drivers = pd.read_sql(fleet_drivers_query, conn)

    print(f"Loaded {len(df_ind_drivers)} independent drivers, {len(df_operators)} operators, and {len(df_fleet_drivers)} fleet-managed drivers.")

    # Write to Excel with custom styling
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df_ind_drivers.to_excel(writer, sheet_name='Independent Driver Logins', index=False)
        df_operators.to_excel(writer, sheet_name='Fleet Operator Logins', index=False)
        df_fleet_drivers.to_excel(writer, sheet_name='Fleet Drivers (No Direct Login)', index=False)
        
        # Summary Sheet
        summary_data = {
            "Metric / Parameter": [
                "Total Independent Driver Accounts (Direct Logins)",
                "Total Fleet Operator Accounts (Operator Logins)",
                "Total Fleet-Managed Drivers (Logins Disabled)",
                "Total Partner Accounts Registered",
                "Portal Web URL",
                "Demo / Sandbox Test OTP",
                "Fleet Driver Access Rule",
                "Export Timestamp"
            ],
            "Value / Rule Details": [
                len(df_ind_drivers),
                len(df_operators),
                len(df_fleet_drivers),
                len(df_ind_drivers) + len(df_operators) + len(df_fleet_drivers),
                "http://localhost:3002/",
                "1234",
                "Drivers under a Fleet Operator MUST NOT have individual logins. All fleet accounting and payments are managed through the Fleet Operator account.",
                datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            ]
        }
        pd.DataFrame(summary_data).to_excel(writer, sheet_name='Summary & Access Rules', index=False)

        # Style Worksheets
        wb = writer.book
        
        # Header Fills
        fill_blue = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_emerald = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
        fill_amber = PatternFill(start_color="9A3412", end_color="9A3412", fill_type="solid")
        fill_slate = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Segoe UI", size=10)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            ws.views.sheetView[0].showGridLines = True
            
            if sheetname == 'Independent Driver Logins':
                fill_to_use = fill_blue
            elif sheetname == 'Fleet Operator Logins':
                fill_to_use = fill_emerald
            elif sheetname == 'Fleet Drivers (No Direct Login)':
                fill_to_use = fill_amber
            else:
                fill_to_use = fill_slate

            for col_idx, col in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
                cell = col[0]
                cell.fill = fill_to_use
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[1].height = 28

            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = regular_font
                    cell.border = thin_border
                    if isinstance(cell.value, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

            # Auto-fit column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    print(f"Successfully generated clean credentials file at: {output_file}")
    return output_file

if __name__ == "__main__":
    generate_credentials_excel()
