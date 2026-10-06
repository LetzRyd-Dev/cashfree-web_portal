import os
import pandas as pd
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.database import engine
from sqlalchemy import text

def generate_comprehensive_credentials_excel():
    downloads_path = os.path.expanduser("~/Downloads")
    output_files = [
        os.path.join(downloads_path, "LetzRyd_All_Profiles_Credentials.xlsx"),
        os.path.join(os.getcwd(), "LetzRyd_All_Profiles_Credentials.xlsx")
    ]
    
    print("Connecting to PostgreSQL database and compiling test profiles and credentials...")
    with engine.connect() as conn:
        # 1. Quick Test Profiles (Curated Core Archetypes for immediate UI testing)
        quick_test_data = [
            {
                "Test Persona": "1. Anurag (Primary Fleet Owner)",
                "Role / Portal": "Operator Portal",
                "Registered Mobile": "9691938866",
                "Test OTP": "1234",
                "Identifier / Code": "OPR-HYD-001 (ID: 1)",
                "Company / Fleet Name": "Anurag & RK Fleet Logistics",
                "Vehicle / Fleet Info": "4 Fleet Cars (Dzire, Tour S, etc.)",
                "Account Type": "Fleet Owner (Multi-Car)",
                "Key Features to Test": "Fleet dashboard, 4 cars breakdown, hisaab settlements, operator ticket creation, manager POC"
            },
            {
                "Test Persona": "2. Saleem (Secondary Fleet Owner)",
                "Role / Portal": "Operator Portal",
                "Registered Mobile": "9848012345",
                "Test OTP": "1234",
                "Identifier / Code": "OPR-HYD-002 (ID: 2)",
                "Company / Fleet Name": "Saleem Fleet Logistics",
                "Vehicle / Fleet Info": "2 Fleet Cars",
                "Account Type": "Fleet Owner (Multi-Car)",
                "Key Features to Test": "Tenant isolation verification (sees strictly 2 cars, 0 cars from Anurag), hisaab summary"
            },
            {
                "Test Persona": "3. Operator 940 (Test Fleet)",
                "Role / Portal": "Operator Portal",
                "Registered Mobile": "9848012348",
                "Test OTP": "1234",
                "Identifier / Code": "OP-501-TEST (ID: 940)",
                "Company / Fleet Name": "Samvreeddhi Mobility Fleet (Test)",
                "Vehicle / Fleet Info": "0 Cars (Isolated Test Fleet)",
                "Account Type": "Fleet Owner (Isolated)",
                "Key Features to Test": "Security isolation: Verify 0 drivers & 0 vehicles leak from Operator 1"
            },
            {
                "Test Persona": "4. Vivek (Active Fleet Driver)",
                "Role / Portal": "Driver Portal",
                "Registered Mobile": "9901484683",
                "Test OTP": "1234",
                "Identifier / Code": "LR-DRV-0001 (ID: 1)",
                "Company / Fleet Name": "Managed by Anurag Fleet",
                "Vehicle / Fleet Info": "KA05AQ7692 (Maruti Dzire CNG)",
                "Account Type": "Fleet Managed Driver",
                "Key Features to Test": "Authentic car display, rental plan ₹1000/day, weekly statements, emergency contact card"
            },
            {
                "Test Persona": "5. Sushant (Driver with Balance)",
                "Role / Portal": "Driver Portal",
                "Registered Mobile": "9140631755",
                "Test OTP": "1234",
                "Identifier / Code": "LR-DRV-0202 (ID: 2)",
                "Company / Fleet Name": "Managed by Anurag Fleet",
                "Vehicle / Fleet Info": "KA05AQ7693 (Maruti Tour S)",
                "Account Type": "Fleet Managed Driver",
                "Key Features to Test": "Driver payout calculation, negative balance (-₹1,034.80), settlement breakdown"
            },
            {
                "Test Persona": "6. Aayush (Fleet Driver)",
                "Role / Portal": "Driver Portal",
                "Registered Mobile": "9930420065",
                "Test OTP": "1234",
                "Identifier / Code": "LR-DRV-0312 (ID: 3)",
                "Company / Fleet Name": "Managed by Anurag Fleet",
                "Vehicle / Fleet Info": "MH01DE1234 (Maruti Dzire)",
                "Account Type": "Fleet Managed Driver",
                "Key Features to Test": "Full weekly statement history, settled payments, cash collection audit"
            },
            {
                "Test Persona": "7. Anurag Driver (Driver Profile)",
                "Role / Portal": "Driver Portal",
                "Registered Mobile": "9866941379",
                "Test OTP": "1234",
                "Identifier / Code": "LR-DRV-0004 (ID: 4)",
                "Company / Fleet Name": "Managed by Anurag Fleet",
                "Vehicle / Fleet Info": "TS09UB5678 (Dzire CNG)",
                "Account Type": "Fleet Managed Driver",
                "Key Features to Test": "Trips goal widget, multi-platform breakdown (Uber, Ola, Rapido)"
            },
            {
                "Test Persona": "8. Mohammed Ali (Direct Driver)",
                "Role / Portal": "Driver Portal",
                "Registered Mobile": "9848012346",
                "Test OTP": "1234",
                "Identifier / Code": "LR-DRV-0418 (ID: 5)",
                "Company / Fleet Name": "Direct LetzRyd Driver",
                "Vehicle / Fleet Info": "TS07UA9999 (Maruti WagonR)",
                "Account Type": "Independent Driver",
                "Key Features to Test": "Independent driver accounting, direct UPI payment flow, Operations Desk manager"
            },
            {
                "Test Persona": "9. Anil Verma (Driver with Due)",
                "Role / Portal": "Driver Portal",
                "Registered Mobile": "9848012347",
                "Test OTP": "1234",
                "Identifier / Code": "LR-DRV-0501 (ID: 6)",
                "Company / Fleet Name": "Direct LetzRyd Driver",
                "Vehicle / Fleet Info": "TS08UB1111 (Maruti Dzire)",
                "Account Type": "Independent Driver",
                "Key Features to Test": "Current due outstanding (₹1,820.50), 'Pay Now' online Cashfree payment simulation"
            },
            {
                "Test Persona": "10. Rajeev (Bangalore Direct)",
                "Role / Portal": "Driver Portal",
                "Registered Mobile": "8147414391",
                "Test OTP": "1234",
                "Identifier / Code": "LETZBLR8147414391 (ID: 439323)",
                "Company / Fleet Name": "LetzRyd Direct Bengaluru",
                "Vehicle / Fleet Info": "Unassigned / Newly Enrolled",
                "Account Type": "Independent Driver",
                "Key Features to Test": "Verify 'No Vehicle Assigned' badge (does NOT show fake KA05AQ7692 car!), clean profile"
            },
            {
                "Test Persona": "11. Prashanth K (Bangalore Direct)",
                "Role / Portal": "Driver Portal",
                "Registered Mobile": "9731753883",
                "Test OTP": "1234",
                "Identifier / Code": "LETZBLR9731753883 (ID: 443135)",
                "Company / Fleet Name": "LetzRyd Direct Bengaluru",
                "Vehicle / Fleet Info": "Allocated Vehicle",
                "Account Type": "Independent Driver",
                "Key Features to Test": "City Hub Manager: Manjunath Gowda (9900088220), emergency contact prompt"
            },
            {
                "Test Persona": "12. Vallery Krishna (Hyderabad Direct)",
                "Role / Portal": "Driver Portal",
                "Registered Mobile": "8985447320",
                "Test OTP": "1234",
                "Identifier / Code": "LETZHYD8985447320 (ID: 438586)",
                "Company / Fleet Name": "LetzRyd Direct Hyderabad",
                "Vehicle / Fleet Info": "Allocated Vehicle",
                "Account Type": "Independent Driver",
                "Key Features to Test": "City Hub Manager: K Ramesh (9848022338), support ticket filing"
            }
        ]
        df_quick_test = pd.DataFrame(quick_test_data)

        # 2. Fleet Operators (All 950 registered operators from PostgreSQL)
        operators_query = text("""
            SELECT 
                o.app_operator_id AS "Operator ID",
                COALESCE(
                    NULLIF(o.operator_code, ''), 
                    'LR-OPR-' || LPAD(o.app_operator_id::text, 4, '0')
                ) AS "Operator Code",
                o.company_name AS "Fleet / Company Name",
                COALESCE(o.contact_person_name, o.company_name) AS "Contact Person",
                o.phone AS "Registered Mobile",
                '1234' AS "Login OTP",
                'Fleet Operator' AS "Account Type",
                COALESCE(o.total_vehicles, 0) AS "Total Fleet Vehicles",
                COALESCE(o.active_vehicles, 0) AS "Active Vehicles",
                COALESCE(o.total_drivers, 0) AS "Total Drivers",
                COALESCE(o.cw_to_collect, 0) AS "Fleet Outstanding Due (₹)",
                COALESCE(o.assigned_manager_name, 'LetzRyd Operations Desk') AS "Assigned Account Manager",
                COALESCE(o.assigned_manager_phone, '9988770011') AS "Manager Phone",
                CASE WHEN o.is_active THEN 'Active (Login Enabled)' ELSE 'Inactive' END AS "Account Status"
            FROM app_operators o
            ORDER BY o.app_operator_id ASC;
        """)
        df_operators = pd.read_sql(operators_query, conn)

        # 3. Independent Drivers (Login Enabled)
        ind_drivers_query = text("""
            SELECT 
                d.app_driver_id AS "Driver ID",
                COALESCE(
                    NULLIF(d.driver_code, ''), 
                    'LR-DRV-' || LPAD(d.app_driver_id::text, 4, '0')
                ) AS "Driver Code",
                d.full_name AS "Driver Name",
                d.phone AS "Registered Mobile",
                '1234' AS "Login OTP",
                COALESCE(d.vehicle_reg_number, 'Unassigned') AS "Assigned Vehicle",
                COALESCE(d.vehicle_make, '') AS "Vehicle Make",
                COALESCE(d.vehicle_model, '') AS "Vehicle Model",
                COALESCE(d.vehicle_daily_rate, 1000) AS "Daily Rental Rate (₹)",
                COALESCE(d.assigned_manager_name, 'LetzRyd Operations Desk') AS "Driver Manager",
                COALESCE(d.assigned_manager_phone, '9988770011') AS "Manager Phone",
                COALESCE(d.cw_to_collect, 0) AS "Current Outstanding (₹)",
                CASE WHEN d.is_active THEN 'Active (Direct Login)' ELSE 'Inactive' END AS "Login Status"
            FROM app_drivers d
            WHERE d.operator_id IS NULL OR d.operator_id = 0
            ORDER BY d.app_driver_id ASC;
        """)
        df_ind_drivers = pd.read_sql(ind_drivers_query, conn)

        # 4. Fleet Managed Drivers (Associated with Fleet Operators)
        fleet_drivers_query = text("""
            SELECT DISTINCT ON (d.app_driver_id)
                d.app_driver_id AS "Driver ID",
                COALESCE(
                    NULLIF(d.driver_code, ''), 
                    'LR-DRV-' || LPAD(d.app_driver_id::text, 4, '0')
                ) AS "Driver Code",
                d.full_name AS "Driver Name",
                d.phone AS "Registered Mobile",
                'Managed via Operator' AS "Login Access",
                COALESCE(o.company_name, 'Fleet Operator #' || d.operator_id) AS "Fleet Operator Name",
                COALESCE(o.operator_code, 'LR-OP-' || d.operator_id) AS "Operator Code",
                COALESCE(d.vehicle_reg_number, 'Fleet Vehicle') AS "Assigned Vehicle",
                COALESCE(d.vehicle_make, '') AS "Vehicle Make",
                COALESCE(d.vehicle_model, '') AS "Vehicle Model",
                COALESCE(d.vehicle_daily_rate, 1000) AS "Daily Rental Rate (₹)",
                COALESCE(d.cw_to_collect, 0) AS "Period Balance (₹)"
            FROM app_drivers d
            LEFT JOIN app_operators o ON d.operator_id = o.app_operator_id
            WHERE d.operator_id IS NOT NULL AND d.operator_id > 0
            ORDER BY d.app_driver_id ASC;
        """)
        df_fleet_drivers = pd.read_sql(fleet_drivers_query, conn)

        # 5. Testing Guide & Instructions
        guide_data = {
            "Step / Area": [
                "1. Access Local Web Portal",
                "2. Login Method",
                "3. Standard Test OTP",
                "4. Testing Fleet Operator Experience",
                "5. Testing Tenant Isolation Security",
                "6. Testing Driver Vehicle & Rental View",
                "7. Testing Unassigned Vehicle Badge",
                "8. Testing Profile & Emergency Contact",
                "9. Testing WhatsApp & Phone Support",
                "10. Testing Support Ticket Creation",
                "11. Testing Online Payment (Cashfree)",
                "12. Database Connection Details"
            ],
            "Instructions & Expected Results": [
                "Open browser and go to http://localhost:3002 (or use standard local port).",
                "Select 'Driver' or 'Fleet Operator', enter the 10-digit Registered Mobile number, and click 'Send OTP'.",
                "Enter '1234' on the OTP screen and click 'Verify OTP' to login instantly.",
                "Login as Anurag (9691938866): View 4 fleet cars, individual driver hisaabs, total debt, and fleet deposit status.",
                "Login as Saleem (9848012345) or Operator 940 (9848012348): Confirm strict isolation (they see 0 cars from Anurag).",
                "Login as Vivek (9901484683): View car KA05AQ7692, Maruti Dzire, daily rate ₹1000/day, weekly earnings & deductions.",
                "Login as Rajeev (8147414391): Notice the clean 'No Vehicle Assigned' badge. No fake car is displayed.",
                "Open Profile Screen: Confirm NO fake 'Priya Kumar' appears. Unset contacts show 'Not Provided — Tap to Add Contact'.",
                "Open Support Screen: Tap 'Chat on WhatsApp'. Verify it opens wa.me/919988770011 (official mobile, not broken landline).",
                "Click '+ New Ticket' on Support Screen. File a ticket and verify it appears with real creator ID and status tracking.",
                "Click 'Pay Now' on driver/operator hisaab. Verify payment modal opens with real outstanding balance.",
                "PostgreSQL Cloud SQL at 35.200.196.113:5432/postgres. Backend FastAPI at http://localhost:8000."
            ]
        }
        df_guide = pd.DataFrame(guide_data)

    print(f"Loaded: {len(df_quick_test)} quick test profiles, {len(df_operators)} operators, {len(df_ind_drivers)} direct drivers, {len(df_fleet_drivers)} fleet drivers.")

    # Write to Excel with custom styling across both target locations
    for target_path in output_files:
        with pd.ExcelWriter(target_path, engine='openpyxl') as writer:
            df_quick_test.to_excel(writer, sheet_name='⭐ Quick Test Profiles', index=False)
            df_guide.to_excel(writer, sheet_name='📖 Testing Guide & Scenarios', index=False)
            df_operators.to_excel(writer, sheet_name='🏢 Fleet Operators (950)', index=False)
            df_ind_drivers.to_excel(writer, sheet_name='🚗 Independent Drivers (478)', index=False)
            df_fleet_drivers.to_excel(writer, sheet_name='📋 Fleet Managed Drivers (2112)', index=False)

            wb = writer.book

            # Palettes
            fill_indigo = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")  # Deep Indigo
            fill_teal = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")    # Teal
            fill_blue = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")    # Cobalt Blue
            fill_emerald = PatternFill(start_color="047857", end_color="047857", fill_type="solid") # Emerald
            fill_amber = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")   # Amber
            
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            bold_font = Font(name="Segoe UI", size=10, bold=True)
            regular_font = Font(name="Segoe UI", size=10)
            thin_border = Border(
                left=Side(style='thin', color='E2E8F0'),
                right=Side(style='thin', color='E2E8F0'),
                top=Side(style='thin', color='E2E8F0'),
                bottom=Side(style='thin', color='E2E8F0')
            )

            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                ws.views.sheetView[0].showGridLines = True
                
                # Freeze top row
                ws.freeze_panes = "A2"

                if 'Quick Test' in sheetname:
                    fill_to_use = fill_indigo
                elif 'Testing Guide' in sheetname:
                    fill_to_use = fill_teal
                elif 'Fleet Operators' in sheetname:
                    fill_to_use = fill_emerald
                elif 'Independent Drivers' in sheetname:
                    fill_to_use = fill_blue
                else:
                    fill_to_use = fill_amber

                # Header styling
                for col in ws.iter_cols(min_row=1, max_row=1):
                    cell = col[0]
                    cell.fill = fill_to_use
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.row_dimensions[1].height = 32

                # Data styling
                for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                    ws.row_dimensions[row_idx].height = 22
                    for cell in row:
                        cell.font = regular_font
                        cell.border = thin_border
                        if isinstance(cell.value, (int, float)):
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                        else:
                            cell.alignment = Alignment(horizontal="left", vertical="center")

                # Auto-fit column widths
                for col in ws.columns:
                    max_len = 0
                    for cell in col:
                        val_str = str(cell.value or '')
                        if len(val_str) > max_len:
                            max_len = len(val_str)
                    col_letter = col[0].column_letter
                    # Cap width to 45 for readable wrapping
                    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

        print(f"Generated clean Excel spreadsheet at: {target_path}")

    return output_files[0]

if __name__ == "__main__":
    generate_comprehensive_credentials_excel()
